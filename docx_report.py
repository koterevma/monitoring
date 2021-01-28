import pandas
import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.shared import RGBColor

doc = Document()

today_date = datetime.date.today().strftime("%d-%m-%y-")
path_file_in = Path(f"data/{today_date}info.csv")
path_file_out = Path(f"reports/{today_date}report.docx")

with open(path_file_in, encoding='utf-8') as file:
    
    cv = pandas.read_csv(file, sep = ';')

    row = cv.shape[0]
    col = cv.shape[1]

    table = doc.add_table(col,2)

    datebegin = cv['Date'][0]
    dateend = cv['Date'][row-1]

    table.cell(0,1).text = datebegin + ' - ' + dateend

    j = 0
    zeros = list()
    ones = list()
    variable = list()#zeros and ones

    for column in cv.columns:
        table.cell(j,0).text = column
        j=j+1
        m = 0
        n = 0
        for i in range(row):
            if str(cv[column][i])=='0':
                m = m+1
                if n!=0:
                    variable.append(cv.columns.get_loc(column))
                    break
                if m==row-1:
                    zeros.append(cv.columns.get_loc(column))  
            if str(cv[column][i])=='1':
                ones.append(cv.columns.get_loc(column))
                n = n+1
                if m!=0:
                    variable.append(cv.columns.get_loc(column))
                    break
                if n==row-1:
                    ones.append(cv.columns.get_loc(column))

    for i in zeros:
        table.cell(i,1).text = '0'
    for i in ones:
        table.cell(i,1).text= '1'
    for i in variable:
        table.cell(i,1).text = 'var'
    print(variable)

doc.save(path_file_out)
